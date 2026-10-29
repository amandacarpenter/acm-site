"""One native Word document from the SAME assembled semantic reading content."""
import base64
import io
import re
import zipfile
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree
from PIL import Image


def render_word(markup, output, title):
    # Native PDF text occasionally contains control glyphs that XML 1.0 cannot
    # encode. Preserve word separation rather than failing the entire catalog.
    markup = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", " ", markup)
    soup = BeautifulSoup(markup, "html.parser")
    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.author = "Perplexity Computer"
    style = doc.styles["Normal"]
    style.font.name = "Calibri"; style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(1)
    style.element.get_or_add_rPr().append(OxmlElement("w:lang"))
    style.element.rPr[-1].set(qn("w:val"), "en-US")
    bookmark_number = 0
    expected_images = len(soup.find_all("img"))

    def bookmark(paragraph, name):
        nonlocal bookmark_number
        bookmark_number += 1
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bookmark_number)); start.set(qn("w:name"), name[:40])
        end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(bookmark_number))
        paragraph._p.append(start); paragraph._p.append(end)

    def inline(node, paragraph, bold=False, italic=False):
        if isinstance(node, NavigableString):
            run = paragraph.add_run(str(node))
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            return
        if not isinstance(node, Tag):
            return
        if node.name in ("p", "div", "blockquote") and paragraph.text.strip():
            paragraph.add_run().add_break()
        if node.name == "br":
            paragraph.add_run().add_break(); return
        if node.name == "img":
            raw = base64.b64decode(node["src"].split(",", 1)[1], validate=True)
            with Image.open(io.BytesIO(raw)) as im:
                ratio = im.height / im.width
            match = re.search(r"\bwidth:\s*([\d.]+)pt", node.get("style", ""))
            width = min(6.5, float(match[1]) / 72 if match else 4.5, 7.8 / ratio)
            picture = paragraph.add_run().add_picture(io.BytesIO(raw), width=Inches(width))
            picture._inline.docPr.set("descr", node.get("alt") or "Image description requires review")
            return
        if node.name == "a" and node.get("href"):
            link = OxmlElement("w:hyperlink"); href = node["href"]
            if href.startswith("#"):
                link.set(qn("w:anchor"), href[1:][:40])
            else:
                link.set(qn("r:id"), paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True))
            run = OxmlElement("w:r"); properties = OxmlElement("w:rPr")
            run_style = OxmlElement("w:rStyle"); run_style.set(qn("w:val"), "Hyperlink")
            properties.append(run_style); run.append(properties)
            text = OxmlElement("w:t"); text.text = node.get_text(); run.append(text)
            link.append(run); paragraph._p.append(link)
            return
        for child in node.children:
            inline(child, paragraph, bold or node.name in ("strong", "b"), italic or node.name in ("em", "i"))

    def block(node, container=doc):
        if isinstance(node, NavigableString):
            if str(node).strip():
                container.add_paragraph(str(node).strip())
            return
        if not isinstance(node, Tag):
            return
        if node.get("id"):
            bookmark(container.add_paragraph(), node["id"])
        if node.name == "table":
            rows = node.find_all("tr")
            if not rows:
                return
            columns = max(sum(int(c.get("colspan", 1)) for c in r.find_all(["th", "td"], recursive=False)) for r in rows)
            if not columns:
                return
            caption = node.find("caption", recursive=False)
            if caption:
                paragraph = container.add_paragraph()
                paragraph.paragraph_format.keep_with_next = True
                inline(caption, paragraph, bold=True)
            table = container.add_table(rows=len(rows), cols=columns)
            table.style = "Table Grid"; table.autofit = False
            for col in table.columns:
                col.width = Inches(6.5 / columns)
            # Respect row/column spans rather than dropping or duplicating cells.
            occupied = set()
            for row_index, row in enumerate(rows):
                col_index = 0
                cells = row.find_all(["th", "td"], recursive=False)
                for source in cells:
                    while (row_index, col_index) in occupied:
                        col_index += 1
                    if col_index >= columns:
                        raise ValueError("Word table structure exceeds its column grid")
                    colspan = min(int(source.get("colspan", 1)), columns - col_index)
                    rowspan = min(int(source.get("rowspan", 1)), len(rows) - row_index)
                    cell = table.cell(row_index, col_index)
                    if rowspan > 1 or colspan > 1:
                        cell = cell.merge(table.cell(row_index + rowspan - 1, col_index + colspan - 1))
                    for r in range(row_index, row_index + rowspan):
                        for c in range(col_index, col_index + colspan):
                            occupied.add((r, c))
                    paragraph = cell.paragraphs[0]
                    for child in source.children:
                        inline(child, paragraph, bold=source.name == "th")
                    col_index += colspan
                if row.find_parent("thead") or (row_index == 0 and any(c.name == "th" for c in cells)):
                    marker = OxmlElement("w:tblHeader")
                    table.rows[row_index]._tr.get_or_add_trPr().append(marker)
            return
        if node.name in ("p", "li", "figcaption", "blockquote", "img") or re.fullmatch("h[1-6]", node.name):
            paragraph = container.add_paragraph()
            if re.fullmatch("h[1-6]", node.name):
                paragraph.style = f"Heading {node.name[1]}"
            elif node.name == "li":
                level = min(3, len(node.find_parents(["ol", "ul"])))
                paragraph.style = ("List Number" if node.parent.name == "ol" else "List Bullet") + (f" {level}" if level > 1 else "")
            if node.name == "li":
                for child in node.children:
                    if isinstance(child, Tag) and child.name in ("ul", "ol"):
                        continue
                    inline(child, paragraph)
                for nested in node.find_all(["ul", "ol"], recursive=False):
                    block(nested, container)
            else:
                inline(node, paragraph)
            return
        for child in node.children:
            block(child, container)

    for node in (soup.body or soup).children:
        block(node)
    doc.save(output)
    with zipfile.ZipFile(output) as archive:
        xml = etree.fromstring(archive.read("word/document.xml"))
        ns = {"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
              "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        drawings = xml.findall(".//wp:docPr", ns)
        if len(drawings) != expected_images or any(not d.get("descr") for d in drawings):
            raise ValueError("Word image/alternative-text preservation check failed")
        bookmarks = {b.get(qn("w:name")) for b in xml.findall(".//w:bookmarkStart", ns)}
        if any(a.get(qn("w:anchor")) not in bookmarks
               for a in xml.findall(".//w:hyperlink", ns) if a.get(qn("w:anchor"))):
            raise ValueError("Word internal-link destination was lost")
