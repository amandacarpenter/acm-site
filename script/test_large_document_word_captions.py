"""Keep calendar/table captions and cell paragraph boundaries in native Word."""
import sys
import tempfile
from pathlib import Path
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "server/pdf_pipelines"))
from large_document_word import render_word

with tempfile.TemporaryDirectory() as directory:
    output = Path(directory)/"captions.docx"
    render_word("""<html><body><h1>Calendar</h1>
    <table><caption>AUGUST 2026</caption><thead><tr><th>Day</th></tr></thead>
    <tbody><tr><td><p>First line.</p><p>Second line.</p></td></tr></tbody></table>
    <table><caption>SEPTEMBER 2026</caption><tr><th>Day</th></tr><tr><td>1</td></tr></table>
    <ul><li><p>One.</p><p>Two.</p></li></ul></body></html>""",output,"Calendar")
    doc=Document(output)
    captions=[p for p in doc.paragraphs if p.text in ("AUGUST 2026","SEPTEMBER 2026")]
    assert len(captions)==2
    assert all(p.paragraph_format.keep_with_next for p in captions)
    assert doc.tables[0].cell(1,0).text=="First line.\nSecond line."
    assert any(p.text=="One.\nTwo." for p in doc.paragraphs)
    assert len(doc.tables)==2
print("PASS: exact captions, keep-with-table, cell/list paragraph boundaries and separate tables.")
